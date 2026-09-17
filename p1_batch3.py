# -*- coding: utf-8 -*-
"""P1 Batch3: method cards + toolbox + learning path"""
import sys, os, json, re
sys.stdout.reconfigure(encoding='utf-8', errors='replace')
import docx

BASE = r"g:\A工作中心\风险处置Commercial risk management"
SITE_DIR = os.path.join(BASE, "07_网站源码", "risk-learning-site")
CONTENT_DIR = os.path.join(SITE_DIR, "content")

def md_to_html(content):
    lines = []
    in_list = False
    for line in content.split('\n'):
        p = line.strip()
        if not p: continue
        if p.startswith("# "):
            if in_list: lines.append('</ul>'); in_list = False
            lines.append(f'<h2>{p[2:]}</h2>')
        elif p.startswith("## "):
            if in_list: lines.append('</ul>'); in_list = False
            lines.append(f'<h3>{p[3:]}</h3>')
        elif p.startswith("### "):
            if in_list: lines.append('</ul>'); in_list = False
            lines.append(f'<h4>{p[4:]}</h4>')
        elif p.startswith(("- ", "* ")):
            if not in_list: lines.append('<ul>'); in_list = True
            lines.append(f'<li>{p[2:]}</li>')
        elif p.startswith(("第一，", "第二，", "第三，", "第四，", "第五，")):
            if in_list: lines.append('</ul>'); in_list = False
            lines.append(f'<p><strong>{p[:3]}</strong>{p[3:]}</p>')
        elif p.startswith("> "):
            if in_list: lines.append('</ul>'); in_list = False
            lines.append(f'<blockquote><p>{p[2:]}</p></blockquote>')
        else:
            if in_list: lines.append('</ul>'); in_list = False
            lines.append(f'<p>{p}</p>')
    if in_list: lines.append('</ul>')
    html = '\n'.join(lines)
    html = html.replace('\\', '\\\\').replace('"', '\u300c').replace('"', '\u300d')
    html = html.replace("'", "\\'").replace('\n', '\\n')
    return html

def docx_to_html(fname, subdir):
    fpath = os.path.join(BASE, subdir, fname)
    if not os.path.exists(fpath):
        print(f"  NOT FOUND: {fpath}")
        return ""
    doc = docx.Document(fpath)
    paras = [p.text for p in doc.paragraphs if p.text.strip()]
    lines = []
    in_list = False
    for p in paras:
        p = p.strip()
        if not p: continue
        if p.startswith(("报告人：", "目录")) and len(p) < 80: continue
        if re.match(r'^第[一二三四五六七八九十]+[章节部分][：:\s]', p):
            if in_list: lines.append('</ul>'); in_list = False
            lines.append(f'<h2>{p}</h2>')
        elif re.match(r'^\d+\.\d+\s', p):
            if in_list: lines.append('</ul>'); in_list = False
            lines.append(f'<h3>{p}</h3>')
        elif p.startswith(("第一步", "第二步", "第三步", "第四步", "第五步")):
            if in_list: lines.append('</ul>'); in_list = False
            parts = p.split("：", 1)
            lines.append(f'<p><strong>{parts[0]}：</strong>{parts[1] if len(parts)>1 else p}</p>')
        elif p.startswith(("核心认知：", "核心原则：", "注意事项：", "使用时机：")):
            if in_list: lines.append('</ul>'); in_list = False
            lines.append(f'<blockquote><p>{p}</p></blockquote>')
        else:
            if in_list: lines.append('</ul>'); in_list = False
            lines.append(f'<p>{p}</p>')
    if in_list: lines.append('</ul>')
    html = '\n'.join(lines)
    html = html.replace('\\', '\\\\').replace('"', '\u300c').replace('"', '\u300d')
    html = html.replace("'", "\\'").replace('\n', '\\n')
    return html

# ═══════════════════════════════════════════
# Part A: 5 Method Cards -> data-tools.js
# ═══════════════════════════════════════════
print("=== Part A: Method Cards ===")
method_files = [
    ("method-01-债务重组六步法.md", "tools-method-01", "债务重组六步法", "进阶", "W7",
     "债务重组六步法：风险诊断、债权人梳理、重组方案设计、谈判推进、执行监督、复盘归档。每一步含操作要点与常见陷阱。"),
    ("method-02-谈判五层心法.md", "tools-method-02", "谈判五层心法", "进阶", "W8",
     "谈判五层心法：准备层（信息收集与底线设定）、开局层（锚定效应）、博弈层（让步策略）、收尾层（锁定成果）、复盘层（经验沉淀）。"),
    ("method-03-债转股结构设计.md", "tools-method-03", "债转股结构设计", "高阶", "W8",
     "债转股结构设计：从债主到股东的身份转换。核心要素包括转股价格、股权比例、治理权安排、退出机制。适用场景与风险提示。"),
    ("method-04-共益债务防火墙.md", "tools-method-04", "共益债务防火墙", "高阶", "W10",
     "共益债务防火墙：在破产重整中保护新增融资的安全。法律优先性、资金封闭管理、用途限定、监督机制四大核心设计。"),
    ("method-05-战时五步法.md", "tools-method-05", "战时五步法", "高阶", "W10",
     "战时五步法：企业突发风险时的快速响应流程。止损隔离、核心保全、利益相关方沟通、处置方案执行、复盘沉淀五步。"),
]

tools_articles = {}
for fname, aid, title, diff, week, summary in method_files:
    fpath = os.path.join(CONTENT_DIR, fname)
    if os.path.exists(fpath):
        with open(fpath, "r", encoding="utf-8") as f:
            content = f.read()
        html = md_to_html(content)
        tools_articles[aid] = {
            "id": aid, "title": title,
            "tags": [diff, week, "方法论卡"],
            "difficulty": diff, "week": week,
            "source": "商业风险处置顾问 2026-09",
            "updated": "2026-09-17", "section": "tools", "subsection": "methods",
            "summary": summary, "relatedIds": [],
            "contentHtml": html
        }
        print(f"  OK: {aid} ({len(html)} chars)")

# ═══════════════════════════════════════════
# Part B: Toolbox 8件套 -> data-tools.js
# ═══════════════════════════════════════════
print("\n=== Part B: Toolbox 8件套 ===")
toolbox_items = [
    # Templates (from 03_实操模板/)
    ("模板-风险处置方案.docx", "03_实操模板", "tools-tpl-01", "标准风险处置方案模板",
     "进阶", "W7", "三套预案版：最优和解/保守处置/兜底诉讼。含使用时机、结构拆解、填写指南。", "模板-风险处置方案.docx"),
    ("模板-风险评估报告.docx", "03_实操模板", "tools-tpl-02", "风险评估报告模板",
     "进阶", "W4", "标准化风险评估报告结构：风险概况、影响范围、等级判定、处置建议。", "模板-风险评估报告.docx"),
    ("模板-结案复盘报告.docx", "03_实操模板", "tools-tpl-03", "项目结案复盘报告模板",
     "进阶", "W13", "复盘四步法：目标回顾、结果对比、原因分析、经验沉淀。", "模板-结案复盘报告.docx"),
    ("模板-隔离止损SOP.docx", "03_实操模板", "tools-tpl-04", "风险隔离与止损操作SOP",
     "进阶", "W10", "五层隔离机制：法律隔离、资产隔离、财务隔离、业务隔离、信息隔离。", "模板-隔离止损SOP.docx"),
    ("模板-谈判话术框架.docx", "03_实操模板", "tools-tpl-05", "谈判沟通话术框架",
     "进阶", "W8", "谈判话术分层：开局话术、 probing话术、让步话术、收尾话术。", "模板-谈判话术框架.docx"),
    # Instruments (from 04_评估工具/)
    ("工具-风险尽调清单.xlsx", "04_评估工具", "tools-inst-01", "企业风险全面尽调清单",
     "入门", "W2", "六维度尽调：工商/涉诉/资产/合同/舆情/资金。可直接打分定级。", "工具-风险尽调清单.xlsx"),
    ("工具-损失测算.xlsx", "04_评估工具", "tools-inst-02", "风险损失测算模板",
     "进阶", "W7", "损失测算三维度：直接损失、间接损失、机会成本。含逾期利息与预期损失计算。", "工具-损失测算.xlsx"),
    ("工具-风险等级评估.xlsx", "04_评估工具", "tools-inst-03", "风险等级评估表",
     "入门", "W3", "概率-影响矩阵评级工具。可直接用于风险定级（高/中/低）。", "工具-风险等级评估.xlsx"),
]

for fname, subdir, aid, title, diff, week, summary, dl_name in toolbox_items:
    if fname.endswith('.xlsx'):
        # xlsx: create guide article instead of reading content
        html = f'<h2>{title} - 使用指南</h2>\\n<p>{summary}</p>\\n<p><strong>使用时机：</strong>在对应学习周的任务中直接使用该工具，结合模板讲解文章理解填写逻辑。</p>\\n<p><strong>文件位置：</strong><a href="../assets/{dl_name}" download>{dl_name}</a></p>'
    else:
        html = docx_to_html(fname, subdir)
    if html:
        # Add download link for xlsx files
        if fname.endswith('.xlsx'):
            html += '\\n<p><strong>\\U0001F4E5 下载：</strong><a href="../assets/' + dl_name + '" download>' + dl_name + '</a></p>'
        tools_articles[aid] = {
            "id": aid, "title": title,
            "tags": [diff, week, "模板" if "tpl" in aid else "工具"],
            "difficulty": diff, "week": week,
            "source": f"{fname} 2026-09",
            "updated": "2026-09-17", "section": "tools",
            "subsection": "templates" if "tpl" in aid else "instruments",
            "summary": summary, "relatedIds": [],
            "contentHtml": html,
            "downloadUrl": f"../assets/{dl_name}" if fname.endswith('.xlsx') else ""
        }
        print(f"  OK: {aid} ({len(html)} chars)")

# ═══════════════════════════════════════════
# Write data-tools.js
# ═══════════════════════════════════════════
print(f"\n=== Writing data-tools.js ({len(tools_articles)} articles) ===")

tools_js = "/* data-tools.js - 实操工具箱数据 */\n"
tools_js += "/* P1 Batch3: 5 模板 + 3 工具 + 5 方法论卡 = 13 篇已入站 */\n\n"
tools_js += "SITE_DATA.tools = {\n"
tools_js += '  subsections: [\n'
tools_js += '    { "id": "templates", "title": "模板讲解", "status": "active" },\n'
tools_js += '    { "id": "instruments", "title": "工具指南", "status": "active" },\n'
tools_js += '    { "id": "methods", "title": "方法论卡", "status": "active" }\n'
tools_js += '  ],\n'
tools_js += "  articles: {\n"

items = list(tools_articles.values())
for i, art in enumerate(items):
    tools_js += f'    "{art["id"]}": {{\n'
    for key in ["id", "title", "tags", "difficulty", "week", "source", "updated", "section", "subsection", "summary", "relatedIds", "contentHtml"]:
        val = json.dumps(art.get(key, ""), ensure_ascii=False)
        tools_js += f'      "{key}": {val},\n'
    dl = art.get("downloadUrl", "")
    tools_js += f'      "downloadUrl": "{dl}"\n'
    tools_js += "    }" + ("," if i < len(items) - 1 else "") + "\n"

tools_js += "  }\n};\n"

tools_path = os.path.join(SITE_DIR, "data-tools.js")
with open(tools_path, "w", encoding="utf-8") as f:
    f.write(tools_js)
print(f"[OK] data-tools.js: {os.path.getsize(tools_path) / 1024:.1f} KB")

# ═══════════════════════════════════════════
# Part C: Learning Path -> data-path.js
# ═══════════════════════════════════════════
print("\n=== Part C: Learning Path ===")
# Read learning plan
lp_path = os.path.join(BASE, "02_学习规划", "风险处置学习方案.docx")
if os.path.exists(lp_path):
    doc = docx.Document(lp_path)
    lp_paras = [p.text for p in doc.paragraphs if p.text.strip()]
    print(f"  Learning plan: {len(lp_paras)} paragraphs")
    
    # Create overview article
    lp_html = docx_to_html("风险处置学习方案.docx", "02_学习规划")
    
    path_js = "/* data-path.js - 学习路径数据 */\n"
    path_js += "/* P1 Batch3: 学习路径总览 + 13 周框架 */\n\n"
    path_js += "SITE_DATA.path = {\n"
    path_js += '  title: "13 周成才路径",\n'
    path_js += '  phases: [\n'
    path_js += '    { "id": "foundation", "title": "筑基期", "weeks": "W1-W5", "color": "#3182ce" },\n'
    path_js += '    { "id": "advanced", "title": "进阶期", "weeks": "W6-W9", "color": "#C89B3C" },\n'
    path_js += '    { "id": "mastery", "title": "成型期", "weeks": "W10-W13", "color": "#38A169" }\n'
    path_js += '  ],\n'
    path_js += '  weeks: {},\n'
    path_js += "  articles: {\n"
    path_js += '    "path-overview": {\n'
    path_js += '      "id": "path-overview",\n'
    path_js += '      "title": "风险处置顾问 · 三个月系统成才方案",\n'
    path_js += '      "tags": ["入门", "W0", "学习路径"],\n'
    path_js += '      "difficulty": "入门",\n'
    path_js += '      "week": "W0",\n'
    path_js += '      "source": "风险处置学习方案 2026-09",\n'
    path_js += '      "updated": "2026-09-17",\n'
    path_js += '      "section": "path",\n'
    path_js += '      "subsection": "overview",\n'
    path_js += '      "summary": "从零具备企业风险排查、风险定级、处置方案撰写、债务纠纷止损、谈判预案、风险复盘全套职业能力。学习定位：偏商业处置、非纯法条、可落地止损、能写方案、能算成本、能给决策。",\n'
    path_js += '      "relatedIds": ["theory-career-01", "theory-book01-01"],\n'
    path_js += f'      "contentHtml": "{lp_html}"\n'
    path_js += '    }\n'
    path_js += "  }\n};\n"
    
    path_path = os.path.join(SITE_DIR, "data-path.js")
    with open(path_path, "w", encoding="utf-8") as f:
        f.write(path_js)
    print(f"[OK] data-path.js: {os.path.getsize(path_path) / 1024:.1f} KB")
else:
    print("  NOT FOUND: learning plan docx")

print("\n=== P1 Batch2+3 Summary ===")
print(f"  data-theory.js: 35 articles (25 book reports + 5 whitepaper + 1 investor + 4 career)")
print(f"  data-tools.js: {len(tools_articles)} articles (5 templates + 3 instruments + 5 methods)")
print(f"  data-path.js: 1 article (overview)")
print(f"  Total: {35 + len(tools_articles) + 1} articles")
